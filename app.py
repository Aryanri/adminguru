import os
import base64
import hashlib
import io
import json
import re
import uuid
import sqlite3
import pandas as pd
from datetime import datetime
from flask import Flask, render_template, redirect, url_for, request, flash, send_file, jsonify, session
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, login_user, logout_user, login_required, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from models import db, User, SchoolProfile, Teacher, Student, Archive, ActivityLog, Announcement, CustomModule, wib_now

app = Flask(__name__)
app.config['SECRET_KEY'] = 'siap_administrasi_guru_2026_secret_key'
database_url = os.environ.get('DATABASE_URL') or os.environ.get('POSTGRES_URL')
if database_url:
    if database_url.startswith("postgres://"):
        database_url = database_url.replace("postgres://", "postgresql+psycopg2://", 1)
    elif database_url.startswith("postgresql://") and "+psycopg2" not in database_url:
        database_url = database_url.replace("postgresql://", "postgresql+psycopg2://", 1)
    if "sslmode" not in database_url and "neon.tech" in database_url:
        separator = "&" if "?" in database_url else "?"
        database_url = database_url + separator + "sslmode=require"
    app.config['SQLALCHEMY_DATABASE_URI'] = database_url
else:
    if os.environ.get('VERCEL') or not os.access(os.path.dirname(__file__) or '.', os.W_OK):
        app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:////tmp/database.db'
    else:
        app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///database.db'

if os.environ.get('VERCEL') or not os.access(os.path.dirname(__file__) or '.', os.W_OK):
    app.config['UPLOAD_FOLDER'] = '/tmp/uploads'
    app.instance_path = '/tmp'
else:
    app.config['UPLOAD_FOLDER'] = 'uploads'

app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

try:
    os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)
except Exception as e:
    print("Upload folder creation notice:", e)

try:
    os.makedirs(app.instance_path, exist_ok=True)
except Exception as e:
    print("Instance folder creation notice:", e)

db.init_app(app)

with app.app_context():
    try:
        db.create_all()
        defaults = [
            ('Analisis Minggu & Hari Efektif Belajar (HEM)', 'hem'),
            ('CP-TP-ATP', 'cp_tp_atp'),
            ('Modul Ajar / RPP', 'modul_ajar'),
            ('LKPD Siswa', 'lkpd'),
            ('Program Semester (Prosem)', 'prosem'),
            ('Program Tahunan (Prota)', 'prota'),
            ('Bank Soal Harian', 'soal_harian'),
            ('Bank Soal STS (Sumatif Tengah Semester)', 'soal_sts'),
            ('Bank Soal SAS (Sumatif Akhir Semester)', 'soal_sas'),
            ('Asesmen Sumatif', 'asesmen_sumatif')
        ]
        for d_name, d_slug in defaults:
            existing = CustomModule.query.filter_by(slug=d_slug).first()
            if not existing:
                cm = CustomModule(nama_modul=d_name, slug=d_slug, aktif=True, is_default=True)
                db.session.add(cm)
        db.session.commit()

        if not User.query.filter_by(username='superadmin').first():
            sa = User(
                username='superadmin',
                email='superadmin@admin.com',
                nip='198001012010011001',
                mapel='Administrator',
                password=generate_password_hash('superadmin123'),
                nama_lengkap='Super Administrator',
                role='superadmin',
                status='active'
            )
            db.session.add(sa)
            db.session.commit()
    except Exception as e:
        print("Init DB error:", e)

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'

@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

def log_activity(user_id, aktivitas):
    try:
        log = ActivityLog(user_id=user_id, aktivitas=aktivitas)
        db.session.add(log)
        db.session.commit()
    except Exception as e:
        print("Log error:", e)

@app.before_request
def check_single_device():
    if current_user.is_authenticated:
        if not session.get('device_session_id'):
            session['device_session_id'] = current_user.current_session_id
        elif current_user.current_session_id and current_user.current_session_id != session.get('device_session_id'):
            logout_user()
            flash('Akun Anda telah login di perangkat lain. Sesi ini diakhiri.', 'danger')
            return redirect(url_for('login'))

@app.route('/')
def index():
    if current_user.is_authenticated:
        return redirect(url_for('dashboard'))
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        user = User.query.filter_by(username=username).first()
        
        if user and check_password_hash(user.password, password):
            if user.status != 'active':
                flash('Akun Anda belum diaktifkan oleh Admin. Silakan konfirmasi via WhatsApp ke 081321984848.', 'warning')
                return redirect(url_for('login'))
            
            device_session_id = str(uuid.uuid4())
            session['device_session_id'] = device_session_id
            user.current_session_id = device_session_id
            db.session.commit()
            
            login_user(user)
            log_activity(user.id, "Login ke sistem dari perangkat baru.")
            flash(f'Selamat datang kembali, {user.nama_lengkap}!', 'success')
            return redirect(url_for('dashboard'))
        else:
            flash('Username atau password salah.', 'danger')
    return render_template('login.html', wa_number='081321984848')

@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        username = request.form.get('username')
        password = request.form.get('password')
        nama_lengkap = request.form.get('nama_lengkap')
        email = request.form.get('email', '')
        nip = request.form.get('nip', '')
        mapel = request.form.get('mapel', '')
        
        if User.query.filter_by(username=username).first():
            flash('Username sudah terdaftar.', 'danger')
            return redirect(url_for('register'))
            
        hashed_password = generate_password_hash(password)
        new_user = User(
            username=username,
            email=email,
            nip=nip,
            mapel=mapel,
            password=hashed_password,
            nama_lengkap=nama_lengkap,
            role='guru',
            status='pending'
        )
        db.session.add(new_user)
        db.session.commit()
        
        wa_url = "https://wa.me/6281321984848?text=Halo%20Admin,%20saya%20telah%20mendaftar%20akun%20guru%20dengan%20nama%20%3A%20" + nama_lengkap + "%20(%username%3A%20" + username + ").%20Mohon%20aktivasi%20akun%20saya."
        return render_template('register_success.html', wa_url=wa_url)
    return render_template('register.html')

@app.route('/logout')
@login_required
def logout():
    log_activity(current_user.id, "Logout dari sistem.")
    logout_user()
    session.pop('device_session_id', None)
    flash('Anda telah berhasil logout.', 'info')
    return """
    <html>
    <head><title>Logging out...</title></head>
    <body style="background:#f8f9fa; display:flex; justify-content:center; align-items:center; height:100vh; font-family:sans-serif;">
    <div style="text-align:center;">
        <h4>Mengakhiri Sesi...</h4>
        <p class="text-muted">Sedang memutuskan koneksi dari Puter.js dan Gurava.</p>
    </div>
    <script src="https://js.puter.com/v2/"></script>
    <script>
        document.addEventListener("DOMContentLoaded", function() {
            if (typeof puter !== 'undefined' && puter.auth.isSignedIn()) {
                puter.auth.signOut();
            }
            setTimeout(() => { window.location.href = '/login'; }, 500);
        });
    </script>
    </body>
    </html>
    """

@app.route('/dashboard', methods=['GET', 'POST'])
@login_required
def dashboard():
    profile = SchoolProfile.query.first()
    if not profile:
        profile = SchoolProfile()
        db.session.add(profile)
        db.session.commit()
        
    if request.method == 'POST':
        form_type = request.form.get('form_type')
        if form_type == 'school_profile':
            profile.nama_sekolah = request.form.get('nama_sekolah')
            profile.alamat = request.form.get('alamat')
            profile.nama_kepsek = request.form.get('nama_kepsek')
            profile.nip_kepsek = request.form.get('nip_kepsek')
            
            file = request.files.get('kop_image')
            if file and file.filename != '':
                filename = secure_filename(file.filename)
                file.save(os.path.join(app.config['UPLOAD_FOLDER'], filename))
                profile.kop_sekolah = filename
                
            db.session.commit()
            log_activity(current_user.id, "Memperbarui Profil Sekolah & Kop Surat.")
            flash('Profil sekolah dan kop surat berhasil diperbarui.', 'success')
        elif form_type == 'user_profile':
            current_user.nama_lengkap = request.form.get('nama_lengkap')
            current_user.nip = request.form.get('nip')
            current_user.email = request.form.get('email')
            current_user.mapel = request.form.get('mapel')
            db.session.commit()
            log_activity(current_user.id, "Memperbarui Identitas Akun Penyusun Dokumen.")
            flash('Identitas akun penyusun berhasil diperbarui.', 'success')
        return redirect(url_for('dashboard'))
        
    user_logs = ActivityLog.query.filter_by(user_id=current_user.id).order_by(ActivityLog.timestamp.desc()).limit(5).all()
    announcement = Announcement.query.filter_by(aktif=True).order_by(Announcement.id.desc()).first()
    
    return render_template('dashboard.html', profile=profile, user_logs=user_logs, announcement=announcement)

# ==========================================================================
# Analisis Minggu & Hari Efektif Belajar (HEM)
# ==========================================================================
HEM_JENJANG = ['SD / MI', 'SMP / MTs', 'SMA / MA', 'SMK / MAK']
# Alias ke model Flash terbaru, supaya analisis tidak mati saat versi tertentu pensiun.
HEM_GEMINI_MODELS = ['gemini-flash-latest', 'gemini-3.5-flash', 'gemini-flash-lite-latest']
HEM_GROQ_MODEL = 'llama-3.3-70b-versatile'
HEM_DOC_PREFIX = 'kalender-'


def hem_tahun_options():
    """Pilihan kalender tahun pelajaran mulai 2025/2026, bertambah otomatis mengikuti tahun berjalan."""
    return [f"{y}/{y + 1}" for y in range(2025, max(2028, wib_now().year) + 1)]


def hem_sistem_hari_label(sistem_hari):
    return '6 Hari Kerja (Senin - Sabtu)' if str(sistem_hari) == '6' else '5 Hari Kerja (Senin - Jumat)'


def hem_document_prefix(jenjang, tahun_pelajaran):
    return HEM_DOC_PREFIX + re.sub(r'[^a-z0-9]+', '-', f"{tahun_pelajaran} {jenjang}".lower()).strip('-')


def hem_find_document(jenjang, tahun_pelajaran):
    """Nama berkas kalender yang tersimpan untuk jenjang + tahun pelajaran ini (bila ada)."""
    folder = app.config['UPLOAD_FOLDER']
    if not os.path.isdir(folder):
        return None
    prefix = hem_document_prefix(jenjang, tahun_pelajaran)
    for name in sorted(os.listdir(folder)):
        if name.startswith(prefix):
            return name
    return None


def hem_store_document(file_storage, jenjang, tahun_pelajaran):
    """Simpan dokumen kalender dari guru, menggantikan dokumen lama untuk jenjang + tahun yang sama."""
    folder = app.config['UPLOAD_FOLDER']
    os.makedirs(folder, exist_ok=True)
    prefix = hem_document_prefix(jenjang, tahun_pelajaran)
    for name in os.listdir(folder):
        if name.startswith(prefix):
            try:
                os.remove(os.path.join(folder, name))
            except OSError:
                pass
    ext = os.path.splitext(secure_filename(file_storage.filename or ''))[1].lower() or '.dat'
    filename = prefix + ext
    file_storage.save(os.path.join(folder, filename))
    return filename


def hem_document_text(filename, raw):
    """Ambil teks kalender dari berkas Word/teks (tanpa pustaka tambahan)."""
    name = (filename or '').lower()
    if name.endswith('.docx'):
        import zipfile
        with zipfile.ZipFile(io.BytesIO(raw)) as berkas_zip:
            xml = berkas_zip.read('word/document.xml').decode('utf-8', 'ignore')
        xml = xml.replace('</w:p>', '\n').replace('</w:tr>', '\n')
        return re.sub(r'<[^>]+>', '', xml).strip()
    if name.endswith(('.txt', '.csv', '.md')):
        return raw.decode('utf-8', 'ignore').strip()
    return ''


def hem_document_mime(filename):
    name = (filename or '').lower()
    if name.endswith('.pdf'):
        return 'application/pdf'
    if name.endswith('.png'):
        return 'image/png'
    if name.endswith(('.jpg', '.jpeg')):
        return 'image/jpeg'
    if name.endswith('.webp'):
        return 'image/webp'
    return 'application/octet-stream'


def hem_prompt(meta, ada_dokumen):
    hari = 'Senin-Jumat' if str(meta['sistem_hari']) == '5' else 'Senin-Sabtu'
    jumlah_hari = '6' if str(meta['sistem_hari']) == '6' else '5'
    sumber = (
        'Dokumen kalender akademik Kemendikdasmen yang disertakan pada pesan ini.'
        if ada_dokumen else
        'Tidak ada dokumen yang disertakan. Gunakan kalender akademik resmi Kemendikdasmen untuk tahun '
        'pelajaran tersebut, dan tuliskan pada catatan bahwa hasilnya wajib diverifikasi dengan dokumen '
        'resmi sekolah.'
    )
    return f"""Bertindaklah sebagai asisten analisis kalender akademik sekolah yang ahli dalam administrasi pendidikan dan perencanaan waktu pembelajaran guru.

KONTEKS SEKOLAH
Nama Sekolah: {meta['nama_sekolah']}
Jenjang: {meta['jenjang']}
Tahun Pelajaran: {meta['tahun_pelajaran']}
Sistem Hari Belajar Sekolah: {meta['sistem_hari_label']} (hari belajar: {hari})

SUMBER DATA
{sumber}

TUGAS ANALISIS
1. Baca dokumen kalender akademik tersebut dan identifikasi: awal dan akhir tahun pelajaran, bulan pembelajaran, minggu dalam setiap bulan, hari belajar sekolah, hari libur (nasional, keagamaan, semester), kegiatan sekolah, asesmen/ujian, serta kegiatan lain yang mempengaruhi pembelajaran. Jika ada data yang tidak jelas, tampilkan sebagai catatan dan jangan berasumsi.
2. Untuk SETIAP BULAN pada tahun pelajaran tersebut hitung:
   - minggu_kalender: jumlah minggu kalender bulan tersebut.
   - minggu_tidak_efektif: minggu libur, minggu tanpa kegiatan pembelajaran, dan minggu kegiatan sekolah yang bukan pembelajaran reguler.
   - hari_sekolah: jumlah hari sekolah menurut sistem {jumlah_hari} hari kerja ({hari}).
   - hari_tidak_efektif: hari libur, hari kegiatan sekolah non pembelajaran, hari asesmen/ujian, dan hari lain yang tidak dipakai untuk pembelajaran.
   - keterangan: agenda singkat bulan tersebut.
3. Jangan memakai kalender umum di luar dokumen/acuan resmi Kemendikdasmen di atas.

KELUARKAN HANYA JSON VALID (tanpa markdown dan tanpa penjelasan tambahan) dengan skema:
{{"bulan": [{{"nama_bulan": "Juli 2025", "minggu_kalender": 5, "minggu_tidak_efektif": 2, "hari_sekolah": 22, "hari_tidak_efektif": 8, "keterangan": "Permulaan tahun ajaran & MPLS"}}], "catatan": ["..."]}}

Semua angka berupa bilangan bulat dan tidak negatif. Urutkan bulan dari awal sampai akhir tahun pelajaran."""


def hem_request_ai(prompt, dokumen_nama, doc_bytes, doc_text):
    """Kirim permintaan analisis ke Gemini (membaca PDF/gambar/teks), Groq sebagai cadangan."""
    import time
    import urllib.request

    errors = []
    gemini_key = os.environ.get('GEMINI_API_KEY')
    groq_key = os.environ.get('GROQ_API_KEY')
    ai_text = ''
    sumber_ai = ''

    if gemini_key:
        parts = [{'text': prompt}]
        if doc_text:
            parts.append({'text': 'ISI DOKUMEN KALENDER AKADEMIK:\n' + doc_text[:60000]})
        elif doc_bytes:
            parts.append({'inline_data': {'mime_type': hem_document_mime(dokumen_nama),
                                          'data': base64.b64encode(doc_bytes).decode('ascii')}})
        payload = {'contents': [{'parts': parts}],
                   'generationConfig': {'temperature': 0.2, 'responseMimeType': 'application/json'}}
        # Layanan AI sesekali membalas 503/429 sesaat, jadi model dicoba bergantian sebelum menyerah.
        for model in HEM_GEMINI_MODELS:
            url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={gemini_key}"
            for percobaan in range(2):
                try:
                    req = urllib.request.Request(url, data=json.dumps(payload).encode('utf-8'),
                                                 headers={'Content-Type': 'application/json'})
                    with urllib.request.urlopen(req, timeout=60) as response:
                        res = json.loads(response.read().decode('utf-8'))
                    kandidat = (res.get('candidates') or [{}])[0]
                    ai_text = kandidat.get('content', {}).get('parts', [{}])[0].get('text', '') or ''
                    if ai_text:
                        sumber_ai = f'Gemini {model}'
                    break
                except Exception as e:
                    errors.append(f'Gemini {model} (percobaan {percobaan + 1}): {e}')
                    if percobaan == 0:
                        time.sleep(2)
            if ai_text:
                break

    if not ai_text and groq_key and (doc_text or not doc_bytes):
        try:
            isi = prompt if not doc_text else prompt + '\n\nISI DOKUMEN KALENDER AKADEMIK:\n' + doc_text[:60000]
            payload = {'model': HEM_GROQ_MODEL, 'temperature': 0.2,
                       'messages': [{'role': 'user', 'content': isi}]}
            req = urllib.request.Request(
                'https://api.groq.com/openai/v1/chat/completions',
                data=json.dumps(payload).encode('utf-8'),
                headers={'Content-Type': 'application/json', 'Authorization': f'Bearer {groq_key}'})
            with urllib.request.urlopen(req, timeout=120) as response:
                res = json.loads(response.read().decode('utf-8'))
            ai_text = (res.get('choices') or [{}])[0].get('message', {}).get('content', '') or ''
            if ai_text:
                sumber_ai = f'Groq {HEM_GROQ_MODEL}'
        except Exception as e:
            errors.append(f'Groq: {e}')

    return ai_text, sumber_ai, errors


def hem_parse_json(ai_text):
    text = (ai_text or '').strip()
    if text.startswith('```'):
        text = re.sub(r'^```[a-zA-Z]*\s*', '', text)
        text = re.sub(r'```\s*$', '', text).strip()
    awal, akhir = text.find('{'), text.rfind('}')
    if awal == -1 or akhir == -1:
        raise ValueError('AI tidak mengembalikan data JSON.')
    return json.loads(text[awal:akhir + 1])


def hem_build_report(ai_data, meta, catatan_tambahan):
    """Rumus minggu/hari efektif dan rekap total dihitung di server agar tidak ada perhitungan ganda."""
    def angka(nilai):
        try:
            return max(0, int(round(float(str(nilai).strip()))))
        except (TypeError, ValueError):
            return 0

    bulan = []
    for item in (ai_data.get('bulan') or []):
        if not isinstance(item, dict):
            continue
        minggu_kalender = angka(item.get('minggu_kalender'))
        minggu_tidak_efektif = min(angka(item.get('minggu_tidak_efektif')), minggu_kalender)
        hari_sekolah = angka(item.get('hari_sekolah'))
        hari_tidak_efektif = min(angka(item.get('hari_tidak_efektif')), hari_sekolah)
        bulan.append({
            'nama_bulan': str(item.get('nama_bulan') or '-').strip(),
            'minggu_kalender': minggu_kalender,
            'minggu_tidak_efektif': minggu_tidak_efektif,
            'minggu_efektif': minggu_kalender - minggu_tidak_efektif,
            'hari_sekolah': hari_sekolah,
            'hari_tidak_efektif': hari_tidak_efektif,
            'hari_efektif': hari_sekolah - hari_tidak_efektif,
            'keterangan': str(item.get('keterangan') or '').strip(),
        })
    if not bulan:
        raise ValueError('AI tidak mengembalikan rincian bulan.')

    kunci = ('minggu_kalender', 'minggu_tidak_efektif', 'minggu_efektif',
             'hari_sekolah', 'hari_tidak_efektif', 'hari_efektif')
    total = {k: sum(b[k] for b in bulan) for k in kunci}
    catatan = [str(c).strip() for c in (ai_data.get('catatan') or []) if str(c).strip()]
    catatan.extend(catatan_tambahan)
    if len(bulan) != 12:
        catatan.append(f'Catatan sistem: hanya {len(bulan)} bulan yang terbaca dari kalender akademik.')
    return {'meta': meta, 'bulan': bulan, 'total': total, 'catatan': catatan}


def hem_latest_report(jenis):
    """Analisis HEM terakhir milik pengguna, bila tersimpan di arsip."""
    arsip = Archive.query.filter_by(user_id=current_user.id, jenis_administrasi=jenis).order_by(Archive.created_at.desc()).first()
    if not arsip:
        return None
    try:
        data = json.loads(arsip.hasil_konten)
    except (TypeError, ValueError):
        return None
    return data if isinstance(data, dict) and data.get('bulan') else None


def hem_page_context(jenis, report=None, error=None, form=None):
    folder = app.config['UPLOAD_FOLDER']
    tersimpan = sorted(n for n in (os.listdir(folder) if os.path.isdir(folder) else []) if n.startswith(HEM_DOC_PREFIX))
    if report is None:
        report = hem_latest_report(jenis)
    tahun_options = hem_tahun_options()
    return {
        'report': report,
        'report_json': json.dumps(report, ensure_ascii=False) if report else '',
        'error': error,
        'jenjang_options': HEM_JENJANG,
        'tahun_options': tahun_options,
        'dokumen_tersimpan': tersimpan,
        'form': form or {'jenjang': HEM_JENJANG[1], 'tahun_pelajaran': tahun_options[-1], 'sistem_hari': '5'},
    }


def hem_analyze_request(req, profile, jenis):
    """Jalankan analisis HEM dari form guru, lalu simpan hasilnya ke arsip."""
    tahun_options = hem_tahun_options()
    jenjang = req.form.get('jenjang', HEM_JENJANG[1])
    if jenjang not in HEM_JENJANG:
        jenjang = HEM_JENJANG[1]
    tahun_pelajaran = req.form.get('tahun_pelajaran', tahun_options[-1])
    if tahun_pelajaran not in tahun_options:
        tahun_pelajaran = tahun_options[-1]
    sistem_hari = '6' if str(req.form.get('sistem_hari')) == '6' else '5'
    form = {'jenjang': jenjang, 'tahun_pelajaran': tahun_pelajaran, 'sistem_hari': sistem_hari}

    unggahan = req.files.get('kalender_file')
    try:
        if unggahan and unggahan.filename:
            dokumen = hem_store_document(unggahan, jenjang, tahun_pelajaran)
            flash(f'Dokumen kalender "{unggahan.filename}" tersimpan dan dipakai untuk analisis ini.', 'info')
        else:
            dokumen = hem_find_document(jenjang, tahun_pelajaran)
    except Exception as e:
        print('HEM upload error:', e)
        return hem_page_context(jenis, error=f'Gagal menyimpan dokumen kalender: {e}', form=form)

    doc_bytes, doc_text = b'', ''
    if dokumen:
        try:
            with open(os.path.join(app.config['UPLOAD_FOLDER'], dokumen), 'rb') as berkas:
                doc_bytes = berkas.read()
            doc_text = hem_document_text(dokumen, doc_bytes)
        except OSError as e:
            doc_bytes, doc_text = b'', ''
            print('HEM read error:', e)

    meta = {
        'nama_sekolah': (profile.nama_sekolah if profile and profile.nama_sekolah else '-'),
        'jenjang': jenjang,
        'tahun_pelajaran': tahun_pelajaran,
        'sistem_hari': sistem_hari,
        'sistem_hari_label': hem_sistem_hari_label(sistem_hari),
        'sumber_dokumen': f'Dokumen kalender: {dokumen}' if dokumen else 'Tanpa dokumen kalender (acuan kalender resmi Kemendikdasmen dari pengetahuan AI)',
        'dibuat': wib_now().strftime('%d-%m-%Y %H:%M'),
    }

    catatan_sistem = []
    if not dokumen:
        catatan_sistem.append('Catatan sistem: analisis ini dibuat tanpa dokumen kalender yang diunggah. Unggah kalender resmi Kemendikdasmen untuk hasil yang pasti.')
    elif doc_bytes and not doc_text:
        catatan_sistem.append('Catatan sistem: dokumen dibaca langsung oleh AI sebagai berkas ' + hem_document_mime(dokumen).split('/')[-1] + '.')

    prompt = hem_prompt(meta, bool(dokumen))
    ai_text, sumber_ai, errors = hem_request_ai(prompt, dokumen, doc_bytes, doc_text)
    if not ai_text:
        pesan = 'Analisis AI gagal. ' + (' | '.join(errors) if errors else 'Tidak ada API Key AI yang tersedia (GEMINI_API_KEY / GROQ_API_KEY).')
        return hem_page_context(jenis, error=pesan, form=form)

    try:
        report = hem_build_report(hem_parse_json(ai_text), meta, catatan_sistem)
    except (ValueError, TypeError) as e:
        return hem_page_context(jenis, error=f'Hasil analisis AI tidak dapat dibaca: {e}', form=form)

    report['meta']['mesin_ai'] = sumber_ai
    isi_arsip = json.dumps(report, ensure_ascii=False)
    try:
        arsip = Archive(
            user_id=current_user.id,
            jenis_administrasi=jenis,
            topik=f"Analisis HEM {tahun_pelajaran} - {jenjang}",
            ruang_lingkup=f"{meta['nama_sekolah']} | {meta['sistem_hari_label']} | {meta['sumber_dokumen']}",
            parameter_hash=hashlib.md5(f"hem-{current_user.id}-{jenjang}-{tahun_pelajaran}-{sistem_hari}-{dokumen}-{len(isi_arsip)}".encode('utf-8')).hexdigest(),
            hasil_konten=isi_arsip,
        )
        db.session.add(arsip)
        db.session.commit()
        log_activity(current_user.id, f"Analisis HEM {tahun_pelajaran} ({jenjang}).")
        flash('Analisis HEM selesai dan tersimpan di arsip.', 'success')
    except Exception as e:
        db.session.rollback()
        print('HEM archive error:', e)
        flash(f'Analisis berhasil, tetapi gagal disimpan ke arsip: {e}', 'warning')

    return hem_page_context(jenis, report=report, form=form)


@app.route('/generator')
@login_required
def generator_index():
    first_mod = CustomModule.query.filter_by(aktif=True).order_by(CustomModule.is_default.desc(), CustomModule.id.asc()).first()
    if not first_mod:
        flash('Tidak ada modul aktif saat ini.', 'danger')
        return redirect(url_for('dashboard'))
    return redirect(url_for('generator', modul_slug=first_mod.slug))

@app.route('/generator/<modul_slug>', methods=['GET', 'POST'])
@login_required
def generator(modul_slug):
    all_modules = CustomModule.query.filter_by(aktif=True).all()
    hasil_generate = None
    
    mod_obj = CustomModule.query.filter_by(slug=modul_slug, aktif=True).first()
    if not mod_obj:
        flash('Modul tidak ditemukan atau sedang dinonaktifkan (gangguan/maintenance).', 'danger')
        first_mod = CustomModule.query.filter_by(aktif=True).first()
        if first_mod:
            return redirect(url_for('generator', modul_slug=first_mod.slug))
        else:
            flash('Tidak ada modul aktif saat ini.', 'danger')
            return redirect(url_for('dashboard'))
            
    jenis = mod_obj.nama_modul
    profile = SchoolProfile.query.first()

    hem_context = hem_page_context(jenis) if modul_slug == 'hem' else {}
    if modul_slug == 'hem' and request.method == 'POST':
        hem_context = hem_analyze_request(request, profile, jenis)
        return render_template('generator.html', all_modules=all_modules, hasil_generate=None, active_slug=modul_slug, current_jenis=jenis, hem_context=hem_context)

    if request.method == 'POST':
        try:
            jenjang = request.form.get('jenjang', 'SMP / MTs')
            kelas_fase_val = request.form.get('kelas_fase', 'Kelas 7 | Fase D (Kelas 7-9)')
            parts = kelas_fase_val.split('|')
            kelas = parts[0].strip() if len(parts) > 0 else 'Kelas 7'
            fase = parts[1].strip() if len(parts) > 1 else 'Fase D (Kelas 7-9)'

            kurikulum = request.form.get('kurikulum', 'Kurikulum Merdeka (Kumer)')
            mapel = request.form.get('mapel', 'Pendidikan Agama Islam dan Budi Pekerti')
            metode = request.form.get('metode', 'Project Based Learning (PjBL)')
            aktivitas = request.form.get('aktivitas', 'Pembelajaran Berbasis Inkuiri')
            kesulitan = request.form.get('kesulitan', 'C3')
            jml_pg = request.form.get('jml_pg', '10')
            jml_esai = request.form.get('jml_esai', '5')
            jml_isian = request.form.get('jml_isian', '5')
            topik = request.form.get('topik', 'Materi Pembelajaran')
            ruang_lingkup = request.form.get('ruang_lingkup', 'Sub-Materi Utama')
            
            param_str = f"AIV3-{jenis}-{jenjang}-{kelas}-{fase}-{kurikulum}-{mapel}-{topik}-{ruang_lingkup}-{kesulitan}-{current_user.id}"
            param_hash = hashlib.md5(param_str.encode('utf-8')).hexdigest()
            
            cached = Archive.query.filter_by(parameter_hash=param_hash).first()
            if cached:
                hasil_generate = cached.hasil_konten
                log_activity(current_user.id, f"Mengambil Arsip Cache untuk {jenis}: {topik}")
                flash('Hasil diambil dari Cache Arsip (Menghemat Token API).', 'info')
            else:
                kop_html = ""
                if profile and profile.kop_sekolah:
                    kop_html = f"<div class='text-center mb-3'><img src='{url_for('uploaded_file', filename=profile.kop_sekolah)}' style='max-width:100%; height:auto; max-height:140px;' alt='Kop Surat'></div><hr style='border: 2px solid #333;'>"
                else:
                    kop_html = f"<div class='text-center mb-3'><h3>{profile.nama_sekolah if profile else 'SEMENCA'}</h3><p>{profile.alamat if profile else ''}</p></div><hr style='border: 2px solid #333;'>"

                guru_nip = current_user.nip if current_user.nip else '-'
                guru_email = current_user.email if current_user.email else '-'
                guru_mapel = current_user.mapel if current_user.mapel else (mapel if mapel else '-')

                if modul_slug in ['soal_harian', 'soal_sts', 'soal_sas', 'asesmen_sumatif']:
                    try:
                        total_soal = int(jml_pg) + int(jml_isian) + int(jml_esai)
                    except:
                        total_soal = 20
                        
                    prompt_ai = f"""Bertindaklah sebagai guru ahli pembuat instrumen evaluasi. Buatkan soal untuk parameter berikut:
Mata Pelajaran: {mapel}
Materi/Topik: {topik}
Ruang Lingkup: {ruang_lingkup}
Jenjang: {jenjang} ({kelas} / {fase})
Kurikulum: {kurikulum}
Level Kognitif: {kesulitan}
Jumlah Soal PG: {jml_pg}
Jumlah Soal Isian Singkat: {jml_isian}
Jumlah Soal Esai: {jml_esai}
Total Keseluruhan Soal: {total_soal} soal

PENTING: Pada bagian "III. KISI-KISI PENULISAN SOAL", Anda WAJIB membuat baris tabel (<tr>) SEBANYAK TOTAL KESELURUHAN SOAL ({total_soal} baris). Setiap soal yang dibuat di bagian naskah soal harus ada pemetaan kisi-kisinya di tabel ini.

Keluarkan HANYA HTML MURNI (tanpa format markdown ```html, jangan gunakan tag body/html). 
Gunakan struktur persis seperti ini:
<div class='mb-1 text-center fs-5'><b>PERANGKAT EVALUASI PEMBELAJARAN</b></div>
<div class='mb-4 text-center text-muted small'>Mata Pelajaran: {mapel} | Materi: {topik} | Jenjang: {jenjang} ({kelas})</div>
<div class='fw-bold mt-2 mb-2 bg-light p-2 border-start border-4 border-primary'>I. NASKAH SOAL EVALUASI</div>
<div class='fw-bold mt-3'>A. PILIHAN GANDA</div>
<div class='ps-2 mb-3'>
  <b>1.</b> [Soal PG 1 yang spesifik dan menantang]<br>
  &nbsp;&nbsp;&nbsp;&nbsp;A. [Opsi]<br>
  &nbsp;&nbsp;&nbsp;&nbsp;B. [Opsi]<br>
  &nbsp;&nbsp;&nbsp;&nbsp;C. [Opsi]<br>
  &nbsp;&nbsp;&nbsp;&nbsp;D. [Opsi]<br><br>
</div>
<div class='fw-bold mt-3'>B. ISIAN SINGKAT</div>
<div class='ps-2 mb-3'>
  <b>1.</b> [Soal Isian 1 yang spesifik] ....................<br><br>
</div>
<div class='fw-bold mt-3'>C. ESAI / URAIAN</div>
<div class='ps-2 mb-4'>
  <b>1.</b> [Soal Esai 1 yang menuntut analisis]<br><br>
</div>
<div class='fw-bold mt-4 mb-2 bg-light p-2 border-start border-4 border-success'>II. KUNCI JAWABAN DAN PEMBAHASAN</div>
<div class='fw-bold mt-3'>A. KUNCI PILIHAN GANDA</div>
<div class='ps-2 mb-3'><b>1.</b> [Kunci] — <i>Pembahasan:</i> [Penjelasan singkat]<br></div>
<div class='fw-bold mt-3'>B. KUNCI ISIAN SINGKAT</div>
<div class='ps-2 mb-3'><b>1.</b> [Jawaban]<br></div>
<div class='fw-bold mt-3'>C. KUNCI ESAI</div>
<div class='ps-2 mb-4'><b>1.</b> <i>Rubrik:</i> [Kriteria jawaban benar]<br></div>
<div class='fw-bold mt-4 mb-2 bg-light p-2 border-start border-4 border-warning'>III. KISI-KISI PENULISAN SOAL</div>
<div class='table-responsive mt-2'>
<table class='table table-bordered align-middle'>
<thead class='table-light text-center'><tr><th>No</th><th>CP / TP</th><th>Materi</th><th>Indikator Soal</th><th>Bentuk Soal</th><th>Level</th></tr></thead>
<tbody>
<tr><td class='text-center'>1</td><td>[Tujuan Pembelajaran]</td><td>{topik}</td><td>[Indikator Soal PG 1]</td><td>PG</td><td class='text-center'>{kesulitan}</td></tr>
<!-- LANJUTKAN MEMBUAT BARIS (<tr>) UNTUK SEMUA SOAL PG, LALU SEMUA SOAL ISIAN, LALU SEMUA SOAL ESAI. PASTIKAN JUMLAH BARIS SAMA PERSIS DENGAN TOTAL SOAL ({total_soal} BARIS). -->
</tbody>
</table>
</div>
"""
                elif modul_slug == 'cp_tp_atp':
                    prompt_ai = f"""Bertindaklah sebagai guru ahli pembuat dokumen administrasi. Buatkan dokumen CP, TP, dan ATP untuk parameter berikut:
Mata Pelajaran: {mapel}
Materi/Topik: {topik}
Ruang Lingkup: {ruang_lingkup}
Jenjang: {jenjang} ({kelas} / {fase})
Kurikulum: {kurikulum}
Pendekatan/Metode: {metode}
Aktivitas: {aktivitas}

Keluarkan HANYA HTML MURNI tanpa markdown ```html, tanpa tag html/body.
Gunakan format tabel dan struktur yang rapi (mirip dokumen resmi):
<h3 class='text-center'>PEMETAAN CAPAIAN PEMBELAJARAN (CP), TUJUAN PEMBELAJARAN (TP)<br>DAN ALUR TUJUAN PEMBELAJARAN (ATP)</h3>
<hr>
<table class='table table-bordered mb-4'>
<tr><th style='width: 25%'>Mata Pelajaran</th><td>{mapel}</td></tr>
<tr><th>Kelas / Fase</th><td>{kelas} / {fase}</td></tr>
<tr><th>Tahun Pelajaran</th><td>2025/2026</td></tr>
</table>
<h4 class='text-primary mt-4'>I. CAPAIAN PEMBELAJARAN (CP) - FASE {fase.split(' ')[0]}</h4>
[Buat tabel Elemen dan Capaian Pembelajaran secara mendalam terkait {topik}]
<h4 class='text-primary mt-4'>II. TUJUAN PEMBELAJARAN (TP) DAN ALUR TUJUAN PEMBELAJARAN (ATP)</h4>
[Buat tabel dengan kolom No, Elemen, Materi/Bab, Tujuan Pembelajaran (TP), Alur Tujuan Pembelajaran (ATP), dan Justifikasi]
"""
                else:
                    prompt_ai = f"""Bertindaklah sebagai guru ahli. Buatkan dokumen administrasi {jenis} untuk parameter berikut:
Mata Pelajaran: {mapel}
Materi/Topik: {topik}
Ruang Lingkup: {ruang_lingkup}
Jenjang: {jenjang} ({kelas} / {fase})
Kurikulum: {kurikulum}
Metode: {metode}
Aktivitas: {aktivitas}

Keluarkan HANYA HTML MURNI tanpa markdown ```html, tanpa tag body/html.
Format dokumen semenarik dan serapi mungkin menyerupai dokumen administrasi guru profesional menggunakan Bootstrap classes jika perlu (table-bordered, dll).
Beri judul dokumen <h3>{jenis.upper()}</h3>
Isi semua bagian dengan konten akademis yang panjang, mendalam, dan spesifik sesuai mata pelajaran {mapel}. Jangan hanya mengulang kata-kata dari prompt atau memberikan template kosong. Buat materi nyata yang siap digunakan guru di kelas untuk topik {topik} dengan ruang lingkup {ruang_lingkup}.
"""

                import urllib.request
                import json
                
                puter_key = os.environ.get("PUTER_API_KEY")
                gemini_key = os.environ.get("GEMINI_API_KEY")
                groq_key = os.environ.get("GROQ_API_KEY")
                
                ai_text = None
                error_msgs = []
                used_api = ""
                
                # 0. PUTER.JS CLIENT-SIDE (Jika berhasil di-generate via JavaScript di frontend)
                ai_generated_content = request.form.get('ai_generated_content')
                if ai_generated_content and ai_generated_content.strip():
                    ai_text = ai_generated_content
                    used_api = "Puter AI (Client-Side)"
                
                # 1. PUTER.JS / PUTER AI (OpenAI Compatible) - Backend Fallback
                if puter_key and not ai_text:
                    try:
                        url = "https://api.puter.com/puterai/openai/v1/chat/completions"
                        payload = {
                            "model": "gpt-4o-mini",
                            "messages": [{"role": "user", "content": prompt_ai}]
                        }
                        req = urllib.request.Request(url, data=json.dumps(payload).encode('utf-8'), headers={
                            'Content-Type': 'application/json',
                            'Authorization': f'Bearer {puter_key}'
                        })
                        with urllib.request.urlopen(req, timeout=30) as response:
                            res_json = json.loads(response.read().decode('utf-8'))
                            ai_text = res_json['choices'][0]['message']['content']
                            used_api = "Puter AI (Server-Side)"
                    except Exception as e:
                        error_msgs.append(f"Puter API Error: {str(e)}")
                        
                # 2. GEMINI AI
                if gemini_key and not ai_text:
                    try:
                        url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-1.5-flash:generateContent?key={gemini_key}"
                        payload = {"contents": [{"parts": [{"text": prompt_ai}]}]}
                        req = urllib.request.Request(url, data=json.dumps(payload).encode('utf-8'), headers={'Content-Type': 'application/json'})
                        with urllib.request.urlopen(req, timeout=30) as response:
                            res_json = json.loads(response.read().decode('utf-8'))
                            ai_text = res_json.get('candidates', [{}])[0].get('content', {}).get('parts', [{}])[0].get('text', '')
                            used_api = "Gemini AI"
                    except Exception as e:
                        error_msgs.append(f"Gemini API Error: {str(e)}")
                        
                # 3. GROQ AI
                if groq_key and not ai_text:
                    try:
                        url = "https://api.groq.com/openai/v1/chat/completions"
                        payload = {
                            "model": "llama-3.3-70b-versatile",
                            "messages": [{"role": "user", "content": prompt_ai}]
                        }
                        req = urllib.request.Request(url, data=json.dumps(payload).encode('utf-8'), headers={
                            'Content-Type': 'application/json',
                            'Authorization': f'Bearer {groq_key}'
                        })
                        with urllib.request.urlopen(req, timeout=30) as response:
                            res_json = json.loads(response.read().decode('utf-8'))
                            ai_text = res_json['choices'][0]['message']['content']
                            used_api = "Groq AI"
                    except Exception as e:
                        error_msgs.append(f"Groq API Error: {str(e)}")
                        
                if ai_text:
                    ai_text = ai_text.strip()
                    if ai_text.startswith("```html"): ai_text = ai_text[7:]
                    if ai_text.startswith("```"): ai_text = ai_text[3:]
                    if ai_text.endswith("```"): ai_text = ai_text[:-3]
                    
                    identitas_html = f"<div class='mb-3 p-3 bg-light border rounded small'><b>IDENTITAS PENYUSUN:</b><br>Nama: {current_user.nama_lengkap}<br>NIP: {guru_nip}<br>Mapel: {mapel}</div>"
                    kepsek_html = f"<div class='row mt-5'><div class='col-6'></div><div class='col-6 text-center'><p>Mengetahui,<br>Kepala Sekolah</p><br><br><br><b>{profile.nama_kepsek if profile else 'Kepala Sekolah'}</b><br>NIP. {profile.nip_kepsek if profile else '-'}</div></div>"
                    
                    if modul_slug in ['soal_harian', 'soal_sts', 'soal_sas', 'asesmen_sumatif']:
                        hasil_generate = kop_html + ai_text
                    else:
                        hasil_generate = kop_html + identitas_html + ai_text + kepsek_html
                        
                else:
                    # ALL FAILED OR NO KEYS
                    hasil_generate = f"{kop_html}" \
                                     f"<div class='alert alert-danger mt-4'>" \
                                     f"<h4><i class='bi bi-exclamation-triangle'></i> API Key Diperlukan atau Terjadi Kesalahan</h4>" \
                                     f"<p>Sistem mencoba menggunakan AI untuk generate modul <b>{jenis}</b>, namun tidak ada API Key yang valid atau semua API mengalami error.</p>" \
                                     f"<p><b>Log Status:</b><br>" + ("<br>".join(error_msgs) if error_msgs else "Tidak ada API Key yang diatur di Environment Variables (Vercel).") + "</p>" \
                                     f"<hr><p class='mb-0'><b>Cara memperbaiki (Tambahkan minimal satu di Vercel Env Vars):</b><br>" \
                                     f"1. <code>PUTER_API_KEY</code> (Untuk Opsi 1 - Puter.js/OpenAI)<br>" \
                                     f"2. <code>GEMINI_API_KEY</code> (Untuk Opsi 2 - Google Gemini AI)<br>" \
                                     f"3. <code>GROQ_API_KEY</code> (Untuk Opsi 3 - Groq Llama/Mixtral)<br>" \
                                     f"Lalu lakukan Redeploy.</p></div>"
                
                new_archive = Archive(
                    user_id=current_user.id,
                    jenis_administrasi=jenis,
                    topik=topik,
                    ruang_lingkup=ruang_lingkup,
                    parameter_hash=param_hash,
                    hasil_konten=hasil_generate
                )
                db.session.add(new_archive)
                db.session.commit()
                log_activity(current_user.id, f"Generate baru administrasi {jenis} ({topik}).")
                flash('Modul berhasil digenerate oleh AI dan disimpan ke arsip.', 'success')
        except Exception as e:
            db.session.rollback()
            print("Generator POST Error:", str(e))
            flash(f'Terjadi kesalahan server saat menyimpan data: {str(e)}', 'danger')
            return redirect(url_for('generator', modul_slug=modul_slug))
            
    return render_template('generator.html', all_modules=all_modules, hasil_generate=hasil_generate, active_slug=modul_slug, current_jenis=jenis, hem_context=hem_context)


@app.route('/uploads/<filename>')
def uploaded_file(filename):
    return send_file(os.path.join(app.config['UPLOAD_FOLDER'], filename))

@app.route('/arsip')
@login_required
def arsip():
    archives = Archive.query.order_by(Archive.created_at.desc()).all()
    return render_template('arsip.html', archives=archives)

@app.route('/arsip/detail/<int:id>')
@login_required
def arsip_detail(id):
    arch = Archive.query.get_or_404(id)
    hem_data = None
    if arch.jenis_administrasi == 'Analisis Minggu & Hari Efektif Belajar (HEM)':
        try:
            data = json.loads(arch.hasil_konten)
            if isinstance(data, dict) and data.get('bulan'):
                hem_data = data
        except (TypeError, ValueError):
            hem_data = None
    return render_template('arsip_detail.html', arch=arch, hem_data=hem_data)

@app.route('/pengaturan', methods=['GET', 'POST'])
@login_required
def pengaturan():
    users = User.query.all() if current_user.role in ['superadmin', 'admin'] else []
    all_logs = ActivityLog.query.order_by(ActivityLog.timestamp.desc()).all() if current_user.role == 'superadmin' else []
    announcement = Announcement.query.filter_by(aktif=True).order_by(Announcement.id.desc()).first()
    
    if request.method == 'POST':
        action = request.form.get('action')
        if action == 'change_password':
            old_pass = request.form.get('old_password')
            new_pass = request.form.get('new_password')
            if check_password_hash(current_user.password, old_pass):
                current_user.password = generate_password_hash(new_pass)
                db.session.commit()
                log_activity(current_user.id, "Mengubah password akun.")
                flash('Password berhasil diubah.', 'success')
            else:
                flash('Password lama salah.', 'danger')
        elif action == 'add_user' and current_user.role in ['superadmin', 'admin']:
            username = request.form.get('username')
            password = request.form.get('password')
            nama_lengkap = request.form.get('nama_lengkap')
            role = request.form.get('role', 'guru')
            nip = request.form.get('nip', '')
            email = request.form.get('email', '')
            mapel = request.form.get('mapel', '')
            
            if User.query.filter_by(username=username).first():
                flash('Username sudah terdaftar.', 'danger')
            else:
                new_user = User(
                    username=username,
                    password=generate_password_hash(password),
                    nama_lengkap=nama_lengkap,
                    role=role,
                    nip=nip,
                    email=email,
                    mapel=mapel,
                    status='active'
                )
                db.session.add(new_user)
                db.session.commit()
                log_activity(current_user.id, f"Menambahkan user baru langsung aktif: {nama_lengkap} ({username})")
                flash(f'User baru {nama_lengkap} berhasil ditambahkan dan langsung aktif.', 'success')
        elif action == 'toggle_announcement' and current_user.role == 'superadmin':
            pesan = request.form.get('pesan_notif')
            aktif_val = True if request.form.get('aktif') == 'on' else False
            if announcement:
                announcement.pesan = pesan
                announcement.aktif = aktif_val
            else:
                new_ann = Announcement(pesan=pesan, aktif=aktif_val)
                db.session.add(new_ann)
            db.session.commit()
            flash('Pengaturan notifikasi/pengumuman diperbarui.', 'success')
        return redirect(url_for('pengaturan'))
        
    return render_template('pengaturan.html', users=users, all_logs=all_logs, announcement=announcement)

@app.route('/admin/user/action/<int:user_id>/<action>')
@login_required
def admin_user_action(user_id, action):
    if current_user.role not in ['superadmin', 'admin']:
        flash('Akses ditolak.', 'danger')
        return redirect(url_for('dashboard'))
        
    target_user = User.query.get_or_404(user_id)
    if action == 'activate':
        target_user.status = 'active'
        flash(f'Akun {target_user.nama_lengkap} berhasil diaktifkan.', 'success')
    elif action == 'delete':
        db.session.delete(target_user)
        flash('Akun berhasil dihapus.', 'success')
    db.session.commit()
    return redirect(url_for('pengaturan'))

@app.route('/superadmin/modul', methods=['GET', 'POST'])
@login_required
def superadmin_modul():
    if current_user.role != 'superadmin':
        flash('Hanya Superadmin yang dapat mengakses menu modul.', 'danger')
        return redirect(url_for('dashboard'))
        
    if request.method == 'POST':
        action = request.form.get('action')
        if action == 'add_module':
            nama_modul = request.form.get('nama_modul')
            slug = nama_modul.lower().replace(' ', '_')
            script_code = request.form.get('script_code')
            file = request.files.get('template_word')
            
            filename = None
            if file and file.filename != '':
                filename = secure_filename(file.filename)
                file.save(os.path.join(app.config['UPLOAD_FOLDER'], filename))
                
            new_mod = CustomModule(
                nama_modul=nama_modul,
                slug=slug,
                template_filename=filename,
                script_code=script_code,
                aktif=True,
                is_default=False
            )
            db.session.add(new_mod)
            db.session.commit()
            log_activity(current_user.id, f"Menambahkan Modul Baru: {nama_modul}")
            flash('Modul baru berhasil ditambahkan ke generator.', 'success')
        return redirect(url_for('superadmin_modul'))
        
    modules = CustomModule.query.order_by(CustomModule.is_default.desc(), CustomModule.id.asc()).all()
    return render_template('superadmin_modul.html', modules=modules)

@app.route('/superadmin/modul/toggle/<int:modul_id>')
@login_required
def superadmin_modul_toggle(modul_id):
    if current_user.role != 'superadmin':
        flash('Akses ditolak.', 'danger')
        return redirect(url_for('dashboard'))
        
    mod = CustomModule.query.get_or_404(modul_id)
    mod.aktif = not mod.aktif
    db.session.commit()
    status_str = "diaktifkan" if mod.aktif else "dinonaktifkan (mode gangguan)"
    flash(f'Modul {mod.nama_modul} berhasil {status_str}.', 'success')
    return redirect(url_for('superadmin_modul'))

if __name__ == '__main__':
    app.run(debug=True, port=5000)

application = app

@app.route('/export/word', methods=['POST'])
@login_required
def export_word():
    html_content = request.form.get('html_content', '')
    title = request.form.get('title', 'Dokumen_Administrasi')
    
    # Wrap in full HTML template for Word interpretation
    full_html = f"""
    <!DOCTYPE html>
    <html>
    <head>
    <meta charset="utf-8">
    <title>{title}</title>
    <style>
        body {{ font-family: 'Arial', sans-serif; font-size: 11pt; line-height: 1.5; color: #333; }}
        table {{ border-collapse: collapse; width: 100%; margin-bottom: 15px; }}
        th, td {{ border: 1px solid #333; padding: 6px 8px; font-size: 10pt; }}
        th {{ background-color: #f2f2f2; text-align: center; }}
        .text-center {{ text-align: center; }}
        .fw-bold {{ font-weight: bold; }}
    </style>
    </head>
    <body>
    {html_content}
    </body>
    </html>
    """
    
    from flask import Response
    response = Response(
        full_html,
        mimetype="application/msword",
        headers={"Content-Disposition": f"attachment;filename={title}.doc"}
    )
    return response

@app.route('/export/excel_hem', methods=['POST'])
@login_required
def export_excel_hem():
    """Unduh hasil analisis HEM sebagai file Excel (.xlsx) siap cetak."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    try:
        report = json.loads(request.form.get('hem_json') or '')
        bulan, total, meta = report['bulan'], report['total'], report.get('meta', {})
    except (TypeError, ValueError, KeyError):
        flash('Data analisis HEM tidak valid. Silakan jalankan analisis ulang.', 'danger')
        return redirect(url_for('generator', modul_slug='hem'))

    BIRU_TUA, ABU_LEMBUT, AKSEN = '1E3D59', 'F1F5F9', 'FFC107'
    tipis = Side(style='thin', color='CBD5E1')
    garis = Border(left=tipis, right=tipis, top=tipis, bottom=tipis)
    fill_biru = PatternFill('solid', fgColor=BIRU_TUA)
    fill_abu = PatternFill('solid', fgColor=ABU_LEMBUT)
    fill_aksen = PatternFill('solid', fgColor=AKSEN)
    font_putih = Font(bold=True, color='FFFFFF', size=11)
    font_aksen = Font(bold=True, color='7C2D12')
    tengah = Alignment(horizontal='center', vertical='center', wrap_text=True)
    kiri = Alignment(horizontal='left', vertical='center', wrap_text=True)

    kolom = 8
    wb = Workbook()
    ws = wb.active
    ws.title = 'Analisis HEM'

    def baris_merge(teks, row, fill=None, font=None, align=tengah, tinggi=20):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=kolom)
        sel = ws.cell(row=row, column=1, value=teks)
        sel.alignment = align
        if fill:
            sel.fill = fill
        if font:
            sel.font = font
        ws.row_dimensions[row].height = tinggi
        return row + 1

    r = 1
    r = baris_merge('LAPORAN ANALISIS MINGGU & HARI EFEKTIF BELAJAR (HEM)', r, fill_biru,
                    Font(bold=True, color='FFFFFF', size=14), tengah, 28)
    r = baris_merge(f"Tahun Pelajaran {meta.get('tahun_pelajaran', '-')} | Jenjang {meta.get('jenjang', '-')} | {meta.get('sistem_hari_label', '-')}",
                    r, fill_aksen, font_aksen, tengah, 20)
    r = baris_merge(f"Nama Sekolah: {meta.get('nama_sekolah', '-')}     |     Sumber data: {meta.get('sumber_dokumen', '-')}",
                    r, fill_abu, Font(italic=True, size=10), kiri, 18)
    r += 1

    kepala = ['Bulan', 'Jumlah Minggu Kalender', 'Minggu Tidak Efektif', 'Minggu Efektif Belajar',
              'Jumlah Hari Sekolah', 'Hari Tidak Efektif', 'Hari Efektif Belajar', 'Keterangan']
    baris_kepala = r
    for i, judul in enumerate(kepala, start=1):
        sel = ws.cell(row=r, column=i, value=judul)
        sel.fill, sel.font, sel.alignment, sel.border = fill_biru, font_putih, tengah, garis
    ws.row_dimensions[r].height = 34
    r += 1

    for urutan, b in enumerate(bulan):
        nilai = [b['nama_bulan'], b['minggu_kalender'], b['minggu_tidak_efektif'], b['minggu_efektif'],
                 b['hari_sekolah'], b['hari_tidak_efektif'], b['hari_efektif'], b['keterangan']]
        for i, val in enumerate(nilai, start=1):
            sel = ws.cell(row=r, column=i, value=val)
            sel.border = garis
            sel.alignment = kiri if i in (1, 8) else tengah
            if urutan % 2 == 1:
                sel.fill = fill_abu
            if i == 1:
                sel.font = Font(bold=True)
            if i in (4, 7):
                sel.fill, sel.font = fill_aksen, font_aksen
        r += 1

    nilai_total = ['TOTAL KESELURUHAN', total['minggu_kalender'], total['minggu_tidak_efektif'], total['minggu_efektif'],
                   total['hari_sekolah'], total['hari_tidak_efektif'], total['hari_efektif'],
                   'Data dasar penyusunan Prota & Prosem']
    for i, val in enumerate(nilai_total, start=1):
        sel = ws.cell(row=r, column=i, value=val)
        sel.border = garis
        sel.alignment = kiri if i in (1, 8) else tengah
        sel.fill, sel.font = (fill_aksen, font_aksen) if i in (4, 7) else (fill_biru, font_putih)
    r += 2

    r = baris_merge('REKAP TOTAL TAHUNAN', r, fill_biru, font_putih, kiri, 22)
    for i, judul in enumerate(('Komponen', 'Jumlah'), start=1):
        sel = ws.cell(row=r, column=i, value=judul)
        sel.fill, sel.font, sel.border = fill_aksen, font_aksen, garis
        sel.alignment = kiri if i == 1 else tengah
    r += 1
    rekap = [
        ('Total Minggu Kalender', total['minggu_kalender']),
        ('Total Minggu Tidak Efektif', total['minggu_tidak_efektif']),
        ('Total Minggu Efektif Belajar', total['minggu_efektif']),
        ('Total Hari Sekolah', total['hari_sekolah']),
        ('Total Hari Tidak Efektif', total['hari_tidak_efektif']),
        ('Total Hari Efektif Belajar', total['hari_efektif']),
    ]
    for urutan, (komponen, jumlah) in enumerate(rekap):
        for i, val in enumerate((komponen, jumlah), start=1):
            sel = ws.cell(row=r, column=i, value=val)
            sel.border = garis
            sel.alignment = kiri if i == 1 else tengah
            if urutan in (2, 5):
                sel.fill, sel.font = fill_aksen, font_aksen
            elif urutan % 2 == 1:
                sel.fill = fill_abu
        r += 1

    if report.get('catatan'):
        r += 1
        r = baris_merge('CATATAN', r, fill_biru, font_putih, kiri, 20)
        for catatan in report['catatan']:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=kolom)
            sel = ws.cell(row=r, column=1, value='- ' + str(catatan))
            sel.alignment = kiri
            r += 1

    for i, lebar in enumerate((16, 15, 15, 15, 15, 15, 15, 46), start=1):
        ws.column_dimensions[get_column_letter(i)].width = lebar
    ws.freeze_panes = f'A{baris_kepala + 1}'
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    aliran = io.BytesIO()
    wb.save(aliran)
    aliran.seek(0)
    nama_file = re.sub(r'[^A-Za-z0-9]+', '_',
                       f"Analisis_HEM_{meta.get('tahun_pelajaran', '')}_{meta.get('jenjang', '')}").strip('_') + '.xlsx'
    return send_file(aliran,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=nama_file)
